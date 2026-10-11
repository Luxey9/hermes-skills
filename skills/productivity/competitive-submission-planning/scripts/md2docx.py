"""md2docx.py — convert a markdown deliverable to a formatted DOCX.

Format is parameterized because every organizer specifies its own. A4 /
Times New Roman 12pt / 1.5 spacing / 3cm margins is a common default for
Indonesian campus and competition documents, but always confirm against the
current rulebook before assuming it.

Usage:
  python md2docx.py source.md out.docx
  python md2docx.py source.md out.docx --font Arial --size 12 \
      --spacing 1.5 --margin-cm 4,2,3.5,2.5

  --margin-cm takes left,right,top,bottom in that order.

Three bugs this script exists to avoid, each of which silently corrupts output:

  1. Code fences. Toggle state ON at the opening fence and consume lines until
     the closing fence. A naive `in_code = not in_code` combined with an inner
     loop that already advanced past the closing fence double-toggles, and the
     remainder of the document gets swallowed into one code block.

  2. Inline emphasis. Split on **bold** AND *italic* in the same pass. A
     splitter that only knows ** leaves stray asterisks throughout the output.

  3. Tables. Consecutive `|` lines must be buffered and rendered as ONE table.
     Rendering per line produces one single-row table per row.

It also re-opens what it wrote and prints the actual margins, font, size and
spacing. Do not skip that step: style assignment can fail silently, and the
only trustworthy check is reading the saved file back.
"""
from __future__ import annotations

import argparse
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PAGE_SIZES = {
    "a4": (21.0, 29.7),
    "letter": (21.59, 27.94),
    "legal": (21.59, 35.56),
}

HEADING_SIZES = {1: 14, 2: 13, 3: 12, 4: 12}


def parse_args(argv=None):
    ap = argparse.ArgumentParser(description="markdown -> DOCX with a fixed format")
    ap.add_argument("source")
    ap.add_argument("output")
    ap.add_argument("--font", default="Times New Roman")
    ap.add_argument("--size", type=float, default=12.0)
    ap.add_argument("--spacing", type=float, default=1.5)
    ap.add_argument("--margin-cm", default="3,3,3,3",
                    help="left,right,top,bottom")
    ap.add_argument("--page", default="a4", choices=sorted(PAGE_SIZES))
    ap.add_argument("--code-font", default="Consolas")
    ap.add_argument("--code-size", type=float, default=9.0)
    ap.add_argument("--table-size", type=float, default=10.0)
    return ap.parse_args(argv)


def configure(doc, args):
    style = doc.styles["Normal"]
    style.font.name = args.font
    style.font.size = Pt(args.size)
    # eastAsia must be set on the rPr or Word substitutes a different face for CJK
    style.element.rPr.rFonts.set(qn("w:eastAsia"), args.font)
    pf = style.paragraph_format
    pf.line_spacing = args.spacing
    pf.space_after = Pt(0)

    left, right, top, bottom = (float(x) for x in args.margin_cm.split(","))
    width, height = PAGE_SIZES[args.page]
    for s in doc.sections:
        s.page_width, s.page_height = Cm(width), Cm(height)
        s.left_margin, s.right_margin = Cm(left), Cm(right)
        s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)


def add_runs(paragraph, text, args, base_bold=False):
    """Render inline **bold** and *italic* into one paragraph."""
    for part in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
            run.bold = base_bold
        run.font.name = args.font


def add_heading(doc, text, level, args):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = args.spacing
    pf.space_before = Pt(12 if level < 3 else 6)
    pf.space_after = Pt(6)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(re.sub(r"\*\*", "", text))
    run.font.name = args.font
    run.font.size = Pt(HEADING_SIZES.get(level, args.size))
    run.bold = True


def add_table(doc, rows, args):
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(width):
            value = row[ci] if ci < len(row) else ""
            cp = cells[ci].paragraphs[0]
            cp.paragraph_format.line_spacing = 1.0
            add_runs(cp, value, args, base_bold=(ri == 0))
            for run in cp.runs:
                run.font.size = Pt(args.table_size)
    doc.add_paragraph()


def flush_table(doc, buffer, args):
    """Render buffered table lines, dropping the |---| separator row."""
    rows = []
    for line in buffer:
        if re.match(r"^\|[\s\-:|]+\|$", line.strip()):
            continue
        rows.append([c.strip() for c in line.strip().strip("|").split("|")])
    add_table(doc, rows, args)


def convert(args):
    doc = Document()
    configure(doc, args)

    lines = open(args.source, encoding="utf-8").read().split("\n")
    i, in_code, table_buffer = 0, False, []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                body = []
                i += 1
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    body.append(lines[i])
                    i += 1
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run("\n".join(body))
                run.font.name = args.code_font
                run.font.size = Pt(args.code_size)
            in_code = False
            i += 1
            continue

        if in_code:
            i += 1
            continue

        if line.startswith("|"):
            table_buffer.append(line)
            nxt = lines[i + 1] if i + 1 < len(lines) else ""
            if not nxt.startswith("|"):
                flush_table(doc, table_buffer, args)
                table_buffer = []
            i += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)", line)
        if heading:
            add_heading(doc, heading.group(2), len(heading.group(1)), args)
            i += 1
            continue

        if line.strip() == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run("_" * 60)
            run.font.name = args.font
            run.font.size = Pt(args.table_size)
            i += 1
            continue

        bullet = re.match(r"^\s*[-*]\s+(.*)", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = args.spacing
            add_runs(p, bullet.group(1), args)
            i += 1
            continue

        numbered = re.match(r"^\s*\d+\.\s+(.*)", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.line_spacing = args.spacing
            add_runs(p, numbered.group(1), args)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = args.spacing
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(p, line.strip(), args)
        i += 1

    doc.save(args.output)
    return args.output


def verify(path, args):
    """Read the saved file back and report what actually landed."""
    doc = Document(path)
    section = doc.sections[0]
    normal = doc.styles["Normal"]
    left, right, top, bottom = (float(x) for x in args.margin_cm.split(","))

    checks = [
        ("margins",
         (round(section.left_margin.cm, 2), round(section.right_margin.cm, 2),
          round(section.top_margin.cm, 2), round(section.bottom_margin.cm, 2)),
         (left, right, top, bottom)),
        ("font", normal.font.name, args.font),
        ("size", normal.font.size.pt, args.size),
        ("spacing", normal.paragraph_format.line_spacing, args.spacing),
    ]

    print(f"saved: {path}")
    print(f"paragraphs: {len(doc.paragraphs)} | tables: {len(doc.tables)}")
    ok = True
    for name, got, want in checks:
        flag = "OK " if got == want else "BAD"
        ok &= got == want
        print(f"{flag} {name}: {got} (want {want})")

    leftovers = [p.text[:80] for p in doc.paragraphs
                 if "**" in p.text or p.text.strip().startswith("|")]
    if leftovers:
        ok = False
        print(f"BAD unconverted markdown in {len(leftovers)} paragraph(s):")
        for item in leftovers[:5]:
            print(f"    {item}")
    else:
        print("OK  no leftover markdown markers")

    return 0 if ok else 1


def main(argv=None):
    args = parse_args(argv)
    path = convert(args)
    return verify(path, args)


if __name__ == "__main__":
    sys.exit(main())
